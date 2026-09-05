#!/usr/bin/env python3
"""
Populate the Banner JV Feeder Template from the "acct-level extract" sheet
produced by epic_gl_extract.py.

Field mapping (extract column -> feeder column), per feeder_map.md, with the
transforms confirmed for this feed:

    rule   -> RULECODE            strip the leading "2": keep only the last 3 digits
    date   -> DOCREFERENCENUMBER  copied as-is
    amt    -> AMOUNT              cents -> dollars (amt / 100)
    type   -> DESCRIPTION         copied as-is
    dbOrCr -> SIGN (+/-)          copied as-is
    bank   -> BANKCODE            strip the trailing "2" (that "2" is the COA, not part of
                                   the bank code) - e.g. "LJ2" -> "LJ", plain 2 -> blank
    fund   -> FUND                copied as-is (blank stays blank)
    org    -> ORG                 copied as-is (blank stays blank)
    acct   -> ACCOUNT             copied as-is
    prog   -> PROGRAM             copied as-is (blank stays blank)

    COA is always 2 (hardcoded, not sourced from the extract).

    ACTIVITY, LOCATION, ENCDNUM, ENCDITEMNUM, ENCDSEQNUM, ENCDACTIONIND, and
    ENCBTYPE are left blank - they are unused for this feed.

Usage:
    python populate_jv_feeder.py \
        --extract July_GL_extract.xlsx \
        --template templates/JVFeederTemplate.xlsx \
        --output July_JV_Feeder.xlsx
"""
from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl

EXTRACT_SHEET_NAME = "acct-level extract"
COA_VALUE = 2

# Columns in the feeder template that this feed never populates.
UNUSED_FEEDER_COLUMNS = [
    "ACTIVITY", "LOCATION", "ENCDNUM", "ENCDITEMNUM",
    "ENCDSEQNUM", "ENCDACTIONIND", "ENCBTYPE",
]


def _strip_leading_rule_prefix(rule) -> object:
    """2125 -> 125, 2301 -> 301. Keeps the last 3 digits of the rule code."""
    if rule is None:
        return None
    s = str(rule).strip()
    return int(s[-3:]) if s[-3:].isdigit() else s[-3:]


def _strip_coa_from_bank(bank) -> object:
    """"LJ2" -> "LJ"; plain 2 (no letters) -> blank. The trailing "2" is the COA,
    which is always 2 for this feed and is written to its own column instead."""
    if bank is None:
        return None
    s = str(bank).strip()
    if s.endswith("2"):
        s = s[:-1]
    return s or None


def read_extract_rows(extract_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(extract_path, data_only=True)
    if EXTRACT_SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"'{EXTRACT_SHEET_NAME}' sheet not found in {extract_path}. "
            f"Available sheets: {wb.sheetnames}"
        )
    ws = wb[EXTRACT_SHEET_NAME]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(header, row)))
    return rows


def map_row(extract_row: dict) -> dict:
    return {
        "RULECODE": _strip_leading_rule_prefix(extract_row.get("rule")),
        "DOCREFERENCENUMBER": extract_row.get("date"),
        "AMOUNT": round(extract_row["amt"] / 100, 2) if extract_row.get("amt") is not None else None,
        "DESCRIPTION": extract_row.get("type"),
        "SIGN (+/-)": extract_row.get("dbOrCr"),
        "BANKCODE": _strip_coa_from_bank(extract_row.get("bank")),
        "COA": COA_VALUE,
        "FUND": extract_row.get("fund"),
        "ORG": extract_row.get("org"),
        "ACCOUNT": extract_row.get("acct"),
        "PROGRAM": extract_row.get("prog"),
        **{col: None for col in UNUSED_FEEDER_COLUMNS},
    }


def write_feeder(feeder_rows: list[dict], template_path: Path, output_path: Path) -> None:
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    missing = [col for col in header if col not in feeder_rows[0]] if feeder_rows else []
    if missing:
        raise ValueError(f"Feeder template has columns with no mapped value: {missing}")

    for row in feeder_rows:
        ws.append([row.get(col) for col in header])

    wb.save(output_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--extract", required=True, type=Path,
                         help="Path to the extract workbook (output of epic_gl_extract.py) "
                              "containing the 'acct-level extract' sheet")
    parser.add_argument("--template", required=True, type=Path,
                         help="Path to the blank JV Feeder Template .xlsx")
    parser.add_argument("--output", required=True, type=Path, help="Path to write the populated feeder .xlsx")
    args = parser.parse_args(argv)

    extract_rows = read_extract_rows(args.extract)
    feeder_rows = [map_row(r) for r in extract_rows]
    write_feeder(feeder_rows, args.template, args.output)

    print(f"Mapped {len(feeder_rows)} rows from '{EXTRACT_SHEET_NAME}' into {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
